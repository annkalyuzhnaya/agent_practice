"""Учебный набор кейса «Документ-инженер». Всё вымышленное, имена файлов начинаются с demo_.

  python tools/demo.py load     положить учебные файлы на шину hub/ (существующие не перезаписываются)
  python tools/demo.py clean    убрать учебные файлы с шины (и всё, что вы собрали из них: в имени есть demo_)
  python tools/demo.py make     пересобрать сам набор demo/hub/ (нужны python-docx и openpyxl; для авторов кейса)

Что в наборе: регламент на 17 разделов (hub/docs/in), его июньская и сентябрьская редакции в хранилище версий,
«правка от коллеги» с двумя лишними изменениями, выгрузка задач и бюджета .xlsx, шаблоны отчёта и сводки,
данные для отчёта, текст правки раздела 3 и три заявки в hub/handoff/.
"""
import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HUB = ROOT / "hub"
DEMO = ROOT / "demo" / "hub"

TASKS = [
    ("Обновить регламент согласования", "А. Иванова (условно)", "выполнено", "2026-09-05", 16),
    ("Свести замечания юристов", "Б. Петров (условно)", "выполнено", "2026-09-08", 6),
    ("Шаблон листа согласования", "А. Иванова (условно)", "выполнено", "2026-09-10", 4),
    ("Обучение согласующих", "В. Сидорова (условно)", "выполнено", "2026-09-12", 12),
    ("Перенос архива документов", "Г. Кузнецов (условно)", "выполнено", "2026-09-15", 20),
    ("Проверка сроков по выборке", "Б. Петров (условно)", "выполнено", "2026-09-17", 8),
    ("Памятка для инициаторов", "В. Сидорова (условно)", "выполнено", "2026-09-19", 5),
    ("Настройка уведомлений", "Г. Кузнецов (условно)", "выполнено", "2026-09-22", 9),
    ("Отчёт по срокам за август", "А. Иванова (условно)", "выполнено", "2026-09-24", 3),
    ("Пилот срочного согласования", "Б. Петров (условно)", "в работе", "2026-10-06", 14),
    ("Реестр типовых документов", "В. Сидорова (условно)", "в работе", "2026-10-09", 7),
    ("Опрос инициаторов", "Г. Кузнецов (условно)", "не начато", "2026-10-15", 0),
]
BUDGET = [("Обучение согласующих", 120, 110), ("Перенос архива", 300, 320), ("Настройка уведомлений", 80, 75), ("Пилот срочного согласования", 150, 60)]
HOURS = [("2026-08-31", 21), ("2026-09-07", 27), ("2026-09-14", 31), ("2026-09-21", 18), ("2026-09-28", 7)]

EDIT = """Срок согласования документа составляет 3 рабочих дня.
Срочные документы согласуются за 1 рабочий день по решению руководителя.
Срок исчисляется со дня, следующего за днём поступления полного комплекта.
"""


def utf8():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


def report_data():
    done = [t for t in TASKS if t[2] == "выполнено"]
    return {
        "направление": "Документооборот",
        "период": "сентябрь 2026",
        "ответственный": "А. Иванова (условно)",
        "выполнено": len(done),
        "всего": len(TASKS),
        "часы_всего": sum(t[4] for t in TASKS),
        "бюджет_план": sum(b[1] for b in BUDGET),
        "бюджет_факт": sum(b[2] for b in BUDGET),
        "задачи": [{"задача": t[0], "исполнитель": t[1], "статус": t[2], "срок": t[3], "часы": t[4]} for t in TASKS],
        "риски": "Пилот срочного согласования сдвинут на октябрь; перенос архива вышел за план по бюджету.",
    }


# ---------- генерация набора (нужны библиотеки) ----------

def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            t.rows[r].cells[c].text = str(value)


def make_reglament(path, edition):
    """edition: 'июнь' (v1), 'сентябрь' (v2), 'коллега' (правка раздела 3 + две лишние правки)."""
    from docx import Document
    june, colleague = edition == "июнь", edition == "коллега"
    doc = Document()
    doc.add_heading("Регламент согласования документов (учебный пример)", 0)
    doc.add_paragraph(f"Редакция: {'июнь' if june else 'сентябрь'} 2026. Организация вымышленная, все сведения условные.")

    doc.add_heading("1. Общие положения", 1)
    doc.add_paragraph("Регламент определяет порядок согласования внутренних документов подразделения.")
    doc.add_paragraph("Регламент обязателен для всех сотрудников, которые готовят, согласуют или утверждают документы.")
    doc.add_paragraph("Согласование ведётся в электронном виде; бумажный лист согласования оформляется по запросу.")

    doc.add_heading("2. Участники и роли", 1)
    doc.add_paragraph("В согласовании участвуют инициатор, согласующие и утверждающий.")
    doc.add_heading("2.1. Инициатор", 2)
    doc.add_paragraph("Инициатор готовит документ, собирает комплект приложений и запускает согласование.")
    doc.add_paragraph("Инициатор отвечает на замечания и вносит правки в согласованные сроки.")
    doc.add_heading("2.2. Согласующие", 2)
    doc.add_paragraph("Состав согласующих зависит от типа документа; базовый состав приведён в таблице.")
    add_table(doc, [["Роль", "Что проверяет", "Срок ответа"],
                    ["Юрист", "Соответствие нормам и договорам", "3 рабочих дня" if colleague else "2 рабочих дня"],
                    ["Финансовый контролёр", "Бюджет и расчёты", "2 рабочих дня"],
                    ["Руководитель направления", "Содержание и сроки", "1 рабочий день"]])
    doc.add_heading("2.3. Утверждающий", 2)
    doc.add_paragraph("Утверждающий принимает итоговое решение после сбора всех согласований.")

    doc.add_heading("3. Сроки согласования", 1)
    if colleague:
        for line in EDIT.strip().splitlines():
            doc.add_paragraph(line)
    else:
        doc.add_paragraph(f"Срок согласования документа составляет {7 if june else 5} рабочих дней.")
        doc.add_paragraph("Срочные документы согласуются за 2 рабочих дня по решению руководителя.")

    doc.add_heading("4. Порядок согласования", 1)
    doc.add_paragraph("Согласование проходит в три шага: подготовка комплекта, параллельное согласование, доработка.")
    doc.add_heading("4.1. Подготовка комплекта", 2)
    doc.add_paragraph("Комплект включает документ, пояснительную записку и расчёты, если они есть.")
    doc.add_paragraph("Неполный комплект возвращается инициатору без рассмотрения.")
    doc.add_heading("4.2. Параллельное согласование", 2)
    doc.add_paragraph("Согласующие получают комплект одновременно и отвечают независимо друг от друга.")
    doc.add_paragraph("Ответ согласующего: «согласовано», «согласовано с замечаниями» или «отклонено».")
    if not june:
        doc.add_heading("4.3. Замечания и доработка", 2)
        doc.add_paragraph("Замечания фиксируются в листе согласования с указанием раздела документа.")
        doc.add_paragraph("После доработки повторно согласуются только изменённые разделы.")

    doc.add_heading("5. Срочные документы", 1)
    doc.add_paragraph("Срочным считается документ, задержка которого ведёт к срыву обязательств подразделения.")
    doc.add_paragraph("Решение о срочности принимает руководитель направления и фиксирует его в листе согласования.")

    doc.add_heading("6. Хранение и версии", 1)
    doc.add_paragraph("Каждая редакция документа хранится отдельной версией с датой и кратким описанием изменений.")
    doc.add_paragraph("Действующую версию определяет утверждающий; прежние версии не удаляются.")

    doc.add_heading("7. Ответственность", 1)
    doc.add_paragraph("Инициатор отвечает за полноту комплекта и достоверность сведений.")
    doc.add_paragraph("Согласующие отвечают за соблюдение сроков ответа" + (" и за качество замечаний." if colleague else "."))

    doc.add_heading("8. Контроль исполнения", 1)
    doc.add_paragraph("Контроль сроков согласования ведёт руководитель направления раз в месяц по выборке документов.")
    doc.add_paragraph("Результаты контроля включаются в ежемесячный отчёт направления.")

    doc.add_heading("9. Заключительные положения", 1)
    doc.add_paragraph("Изменения в регламент вносятся по той же процедуре согласования.")
    doc.add_paragraph("Регламент пересматривается не реже одного раза в год.")

    doc.add_heading("Приложение А. Лист согласования", 1)
    doc.add_paragraph("Форма листа согласования (заполняется на каждый документ).")
    add_table(doc, [["Согласующий", "Решение", "Дата", "Замечания (раздел)"], ["", "", "", ""], ["", "", "", ""]])
    if not june:
        doc.add_heading("Приложение Б. Типы документов", 1)
        doc.add_paragraph("Перечень типов документов и состав согласующих.")
        add_table(doc, [["Тип документа", "Согласующие"],
                        ["Регламент", "Юрист, руководитель направления"],
                        ["Договор", "Юрист, финансовый контролёр"],
                        ["Отчёт", "Руководитель направления"]])
    doc.save(str(path))


def make_report_template(path):
    from docx import Document
    doc = Document()
    doc.add_heading("Отчёт направления «{{направление}}» за {{период}}", 1)
    doc.add_paragraph("Ответственный: {{ответственный}}.")
    doc.add_heading("1. Итоги", 2)
    doc.add_paragraph("Выполнено задач: {{выполнено}} из {{всего}}. Трудозатраты: {{часы_всего}} ч.")
    doc.add_heading("2. Задачи", 2)
    add_table(doc, [["Задача", "Исполнитель", "Статус", "Срок"],
                    ["{{rows:задачи}}{{задача}}", "{{исполнитель}}", "{{статус}}", "{{срок}}"]])
    doc.add_heading("3. Бюджет", 2)
    doc.add_paragraph("План: {{бюджет_план}} тыс. руб., факт: {{бюджет_факт}} тыс. руб.")
    doc.add_heading("4. Риски", 2)
    doc.add_paragraph("{{риски}}")
    doc.save(str(path))


def make_export(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Задачи"
    ws.append(["Задача", "Исполнитель", "Статус", "Срок", "Часы"])
    for t in TASKS:
        ws.append(list(t))
    ws = wb.create_sheet("Бюджет")
    ws.append(["Статья", "План, тыс. руб.", "Факт, тыс. руб."])
    for b in BUDGET:
        ws.append(list(b))
    ws.append(["Итого", f"=SUM(B2:B{len(BUDGET) + 1})", f"=SUM(C2:C{len(BUDGET) + 1})"])
    ws = wb.create_sheet("Часы")
    ws.append(["Неделя с", "Часы"])
    for h in HOURS:
        ws.append(list(h))
    wb.save(str(path))


def make_summary_template(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Свод"
    ws.append(["Сводка за {{период}}"])
    ws.append(["Задача", "Статус", "Часы"])
    ws.append(["{{rows:задачи}}{{задача}}", "{{статус}}", "{{часы}}"])
    ws = wb.create_sheet("Справка")
    ws.append(["Направление", "{{направление}}"])
    ws.append(["Ответственный", "{{ответственный}}"])
    wb.save(str(path))


def request(path, source, title, what, files, attention):
    lines = ["---", "to: docs", "needs: code", "status: new", "stage: принято", f"from: {source}", "date: " + path.name[5:15], "---",
             f"# Заявка: {title}", "", f"**Что сделать.** {what}", "", "**Файлы.** " + "; ".join(f"`{f}`" for f in files), "",
             f"**Обратить внимание.** {attention}", "", f"**Источник.** {source}, {path.name[5:15]} (учебный пример)", ""]
    path.write_text("\n".join(lines), encoding="utf-8", newline="\n")


def cmd_make():
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import versions
    if DEMO.exists():
        shutil.rmtree(DEMO)
    d_in, d_tpl, d_store, d_hand = DEMO / "docs" / "in", DEMO / "docs" / "templates", DEMO / "docs" / "store", DEMO / "handoff"
    for folder in (d_in, d_tpl, d_store, d_hand):
        folder.mkdir(parents=True)
    june = DEMO / "demo_регламент_июнь.docx"
    make_reglament(june, "июнь")
    make_reglament(d_in / "demo_регламент.docx", "сентябрь")
    make_reglament(d_in / "demo_регламент_правка_от_коллеги.docx", "коллега")
    versions.add_version(june, "demo_reglament", "регламент", "июньская редакция (учебный пример)", "2026-06-15T10:00:00", d_store)
    versions.add_version(d_in / "demo_регламент.docx", "demo_reglament", "регламент",
                         "сентябрьская редакция: срок 5 дней, раздел 4.3, приложение Б (учебный пример)", "2026-09-20T10:00:00", d_store)
    june.unlink()
    make_export(d_in / "demo_выгрузка.xlsx")
    (d_in / "demo_данные_отчёта.json").write_text(json.dumps(report_data(), ensure_ascii=False, indent=1), encoding="utf-8")
    (d_in / "demo_правка_сроков.md").write_text(EDIT, encoding="utf-8", newline="\n")
    make_report_template(d_tpl / "demo_шаблон_отчёта.docx")
    make_summary_template(d_tpl / "demo_шаблон_сводки.xlsx")
    request(d_hand / "demo_2026-10-01_правка-сроков-регламента.md", "чат Cowork",
            "Обновить раздел 3 в регламенте согласования",
            "Заменить текст раздела «3. Сроки согласования» на текст из файла правки. Остальные разделы не трогать.",
            ["hub/docs/in/demo_регламент.docx", "hub/docs/in/demo_правка_сроков.md"],
            "Нужен отчёт целостности: изменён только раздел 3. Новую редакцию сохранить версией документа demo_reglament.")
    request(d_hand / "demo_2026-10-01_отчёт-за-сентябрь.md", "страница «Документы»",
            "Собрать отчёт направления за сентябрь по шаблону",
            "Собрать отчёт по шаблону из данных; числа сверить с выгрузкой. Дополнительно — сводку Excel по шаблону сводки.",
            ["hub/docs/templates/demo_шаблон_отчёта.docx", "hub/docs/templates/demo_шаблон_сводки.xlsx",
             "hub/docs/in/demo_данные_отчёта.json", "hub/docs/in/demo_выгрузка.xlsx"],
            "Незаполненных меток быть не должно. Каждое число отчёта должно находиться в выгрузке или данных.")
    request(d_hand / "demo_2026-10-02_проверить-правку-коллеги.md", "telegram",
            "Проверить правку регламента от коллеги",
            "Коллега прислал свою редакцию регламента и говорит, что менял только раздел 3. Проверить, так ли это.",
            ["hub/docs/in/demo_регламент.docx", "hub/docs/in/demo_регламент_правка_от_коллеги.docx"],
            "Если задето что-то кроме раздела 3 — версию не сохранять, перечислить задетые разделы.")
    files = sorted(p.relative_to(DEMO).as_posix() for p in DEMO.rglob("*") if p.is_file())
    print(f"Учебный набор пересобран: demo/hub/, файлов {len(files)}")
    for name in files:
        print("  " + name)


# ---------- загрузка и уборка (только стандартная библиотека) ----------

def cmd_load():
    if not HUB.exists():
        sys.exit("Нет папки hub/ — сначала: python tools/setup.py")
    if not DEMO.exists():
        sys.exit("Нет набора demo/hub/ — пересоберите: python tools/demo.py make")
    copied = skipped = 0
    for src in sorted(DEMO.rglob("*")):
        if src.is_dir():
            continue
        dst = HUB / src.relative_to(DEMO)
        if dst.exists():
            skipped += 1
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, dst)
        copied += 1
    print(f"Учебные файлы на шине: добавлено {copied}" + (f", уже были {skipped}" if skipped else ""))
    print("  hub/docs/in/         demo_регламент.docx, demo_регламент_правка_от_коллеги.docx, demo_выгрузка.xlsx, данные и текст правки")
    print("  hub/docs/templates/  demo_шаблон_отчёта.docx, demo_шаблон_сводки.xlsx")
    print("  hub/docs/store/      demo_reglament: v1 (июнь), v2 (сентябрь); действующая не назначена")
    print("  hub/handoff/         три заявки demo_…")
    print("Скажите Claude: «обработай заявки» — или пройдите задание L2 из LEVELS.md.")


def cmd_clean():
    removed = 0
    for sub in ("handoff", "docs/in", "docs/out", "docs/templates", "docs/checks", "inbox/new"):
        folder = HUB / sub
        for path in folder.glob("*") if folder.exists() else []:
            if path.is_file() and "demo_" in path.name:
                path.unlink()
                removed += 1
    store = HUB / "docs" / "store"
    for folder in store.glob("demo_*") if store.exists() else []:
        if folder.is_dir():
            shutil.rmtree(folder)
            removed += 1
    print(f"Учебные файлы убраны: {removed}. Ваши документы не тронуты.")
    print("Если учебные записи остались в hub/docs/registry.md — обновите: python tools/registry.py export")


def main():
    utf8()
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "load":
        cmd_load()
    elif cmd == "clean":
        cmd_clean()
    elif cmd == "make":
        cmd_make()
    else:
        print(__doc__)


if __name__ == "__main__":
    main()
