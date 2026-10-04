"""Эвалы целостности (уровень L5): «подложи лишнюю правку — валидатор должен поймать».

  python tools/evals.py            прогнать все случаи, напечатать таблицу
  python tools/evals.py --keep     не удалять временную папку (посмотреть файлы)

Каждый случай строит во временной папке пару «до/после», запускает настоящий скрипт и сверяет код возврата
и ключевую фразу с ожиданием. Шину hub/ не трогает. Код возврата: 0 — все случаи прошли, 1 — есть провал.
Свой случай добавляется функцией case_...() в список CASES.
"""
import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import doclib

TOOLS = Path(__file__).resolve().parent


def run(script, *args):
    proc = subprocess.run([sys.executable, str(TOOLS / script), *map(str, args)], capture_output=True, text=True,
                          encoding="utf-8", errors="replace", env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"})
    return proc.returncode, proc.stdout + proc.stderr


def base_doc(path):
    from docx import Document
    doc = Document()
    doc.add_heading("1. Общие положения", 1)
    doc.add_paragraph("Регламент описывает согласование документов (учебный пример).")
    doc.add_heading("2. Участники", 1)
    doc.add_paragraph("В согласовании участвуют инициатор и согласующие.")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Роль", "Срок, дней"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Юрист", "2"
    doc.add_heading("3. Сроки согласования", 1)
    doc.add_paragraph("Срок согласования документа составляет 5 рабочих дней.")
    doc.add_heading("4. Ответственность", 1)
    doc.add_paragraph("Инициатор отвечает за полноту комплекта.")
    doc.save(str(path))


def edit(src, dst, fn):
    from docx import Document
    doc = Document(str(src))
    fn(doc)
    doc.save(str(dst))


def set_text(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            p.runs[0].text = p.text.replace(old, new)
            for r in p.runs[1:]:
                r.text = ""
            return
    raise SystemExit(f"эвал: не нашёл абзац «{old}»")


def allowed_edit(tmp):
    """Честная правка раздела 3 генератором."""
    new = tmp / "new_section.md"
    new.write_text("Срок согласования документа составляет 3 рабочих дня.\n", encoding="utf-8")
    code, out = run("generator.py", "update-section", tmp / "before.docx", "--section", "Сроки", "--text-file", new,
                    "--out", tmp / "after_ok.docx")
    if code:
        raise SystemExit("эвал: генератор не смог сделать базовую правку:\n" + out)
    return tmp / "after_ok.docx"


def case_allowed(tmp):
    return run("integrity_diff.py", tmp / "before.docx", allowed_edit(tmp), "--allow", "Сроки")


def case_extra_text(tmp):
    edit(allowed_edit(tmp), tmp / "after_extra.docx", lambda d: set_text(d, "полноту комплекта", "полноту комплекта и сроки"))
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_extra.docx", "--allow", "Сроки")


def case_table_cell(tmp):
    def fn(doc):
        doc.tables[0].rows[1].cells[1].text = "4"
    edit(allowed_edit(tmp), tmp / "after_table.docx", fn)
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_table.docx", "--allow", "Сроки")


def case_removed(tmp):
    def fn(doc):
        for p in list(doc.paragraphs):
            if "Ответственность" in p.text or "полноту комплекта" in p.text:
                p._p.getparent().remove(p._p)
    edit(allowed_edit(tmp), tmp / "after_removed.docx", fn)
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_removed.docx", "--allow", "Сроки")


def case_renamed(tmp):
    edit(allowed_edit(tmp), tmp / "after_renamed.docx", lambda d: set_text(d, "2. Участники", "2. Участники процесса"))
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_renamed.docx", "--allow", "Сроки")


def bold(doc):
    for p in doc.paragraphs:
        if "полноту комплекта" in p.text:
            p.runs[0].bold = True


def case_format_strict(tmp):
    edit(allowed_edit(tmp), tmp / "after_bold.docx", bold)
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_bold.docx", "--allow", "Сроки")


def case_format_text(tmp):
    edit(allowed_edit(tmp), tmp / "after_bold2.docx", bold)
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_bold2.docx", "--allow", "Сроки", "--text")


def case_not_applied(tmp):
    shutil.copy2(tmp / "before.docx", tmp / "after_same.docx")
    return run("integrity_diff.py", tmp / "before.docx", tmp / "after_same.docx", "--allow", "Сроки")


def case_xlsx_sheet(tmp):
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active.title = "Свод"
    wb.active.append(["Показатель", "Значение"])
    wb.active.append(["Выполнено", 9])
    ws = wb.create_sheet("Бюджет")
    ws.append(["Статья", "Факт"])
    ws.append(["Обучение", 410])
    wb.save(str(tmp / "before.xlsx"))
    wb = load_workbook(str(tmp / "before.xlsx"))
    wb["Свод"]["B2"] = 10
    wb["Бюджет"]["B2"] = 400
    wb.save(str(tmp / "after.xlsx"))
    return run("integrity_diff.py", tmp / "before.xlsx", tmp / "after.xlsx", "--allow", "Свод")


def case_missing_key(tmp):
    from docx import Document
    doc = Document()
    doc.add_heading("Отчёт за {{период}}", 1)
    doc.add_paragraph("Ответственный: {{ответственный}}.")
    doc.save(str(tmp / "tpl.docx"))
    (tmp / "data.json").write_text(json.dumps({"период": "сентябрь 2026"}, ensure_ascii=False), encoding="utf-8")
    return run("generator.py", "fill", tmp / "tpl.docx", tmp / "data.json", "--out", tmp / "filled.docx", "--strict")


def case_invented_number(tmp):
    from docx import Document
    doc = Document()
    doc.add_heading("Отчёт", 1)
    doc.add_paragraph("Выполнено задач: 9 из 12. Экономия составила 37 процентов.")
    doc.save(str(tmp / "report.docx"))
    (tmp / "src.json").write_text(json.dumps({"выполнено": 9, "всего": 12}, ensure_ascii=False), encoding="utf-8")
    return run("numcheck.py", tmp / "report.docx", tmp / "src.json")


# (название, функция, ожидаемый код, фраза, которая должна быть в выводе)
CASES = [
    ("разрешённая правка раздела проходит", case_allowed, 0, "целостность соблюдена"),
    ("лишняя правка текста в другом разделе", case_extra_text, 1, "НАРУШЕНИЕ изменён: 4. Ответственность"),
    ("подменена ячейка таблицы в другом разделе", case_table_cell, 1, "НАРУШЕНИЕ изменён: 2. Участники"),
    ("удалён посторонний раздел", case_removed, 1, "НАРУШЕНИЕ удалён: 4. Ответственность"),
    ("переименован заголовок другого раздела", case_renamed, 1, "НАРУШЕНИЕ"),
    ("изменено только оформление (строгий режим)", case_format_strict, 1, "НАРУШЕНИЕ изменён: 4. Ответственность"),
    ("изменено только оформление (режим --text)", case_format_text, 0, "целостность соблюдена"),
    ("правка не применилась — предупреждение", case_not_applied, 0, "ВНИМАНИЕ"),
    ("Excel: задет лист вне разрешённых", case_xlsx_sheet, 1, "НАРУШЕНИЕ изменён: Бюджет"),
    ("шаблон: нет данных для метки (--strict)", case_missing_key, 2, "НЕ ЗАПОЛНЕНО"),
    ("в отчёте число, которого нет в источнике", case_invented_number, 1, "НЕТ В ИСТОЧНИКАХ: 37"),
]


def main():
    doclib.utf8()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--keep", action="store_true")
    ap.add_argument("--verbose", action="store_true", help="печатать вывод скриптов")
    args = ap.parse_args()
    tmp = Path(tempfile.mkdtemp(prefix="docs_evals_"))
    base_doc(tmp / "before.docx")
    failed = 0
    print(f"{'Случай':<48} {'ждали':>5} {'вышло':>5}  итог")
    for name, fn, want_code, phrase in CASES:
        code, out = fn(tmp)
        ok = code == want_code and phrase in out
        failed += not ok
        print(f"{name:<48} {want_code:>5} {code:>5}  {'ПРОШЁЛ' if ok else 'ПРОВАЛ'}")
        if args.verbose or not ok:
            print("    " + out.strip().replace("\n", "\n    "))
            if not ok:
                print(f"    ожидалась фраза: «{phrase}»")
    if args.keep:
        print(f"Файлы оставлены: {tmp}")
    else:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f"ИТОГ: прошло {len(CASES) - failed} из {len(CASES)}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
