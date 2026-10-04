"""Общая библиотека стенда: разбор docx/xlsx в типизированную структуру с хэшами разделов.

Раздел docx = заголовок + всё до следующего заголовка; раздел xlsx = лист.
У каждого раздела два хэша: text_hash (содержание) и xml_hash (содержание + оформление).
"""
import hashlib
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
HUB = ROOT / "hub"
PREAMBLE = "(преамбула)"

_RSID = re.compile(r'\sw:rsid\w*="[^"]*"')
_HEADING = re.compile(r"(?:Heading|Заголовок)\s*(\d+)", re.I)


def utf8():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


def sha(text):
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def file_sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()[:16]


def _xml(el):
    return _RSID.sub("", el.xml)


def heading_level(p):
    name = p.style.name if p.style is not None and p.style.name else ""
    m = _HEADING.match(name)
    return int(m.group(1)) if m else None


def iter_blocks(doc):
    """Абзацы и таблицы тела документа в порядке следования."""
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(el, doc)
        elif tag == "tbl":
            yield Table(el, doc)


def _new_section(title, level, path):
    return {"title": title, "level": level, "path": path, "paragraphs": [], "tables": [], "_xml": []}


def parse_docx(path):
    doc = Document(str(path))
    cur = _new_section(PREAMBLE, 0, PREAMBLE)
    sections, stack = [cur], []
    for b in iter_blocks(doc):
        if isinstance(b, Table):
            cur["tables"].append([[c.text for c in r.cells] for r in b.rows])
            cur["_xml"].append(_xml(b._tbl))
            continue
        lvl = heading_level(b)
        if lvl and b.text.strip():
            title = b.text.strip()
            stack = [s for s in stack if s[0] < lvl] + [(lvl, title)]
            cur = _new_section(title, lvl, " > ".join(t for _, t in stack))
            sections.append(cur)
        elif b.text.strip():
            cur["paragraphs"].append(b.text)
        cur["_xml"].append(_xml(b._p))
    for s in sections:
        s["text_hash"] = sha(json.dumps([s["title"], s["paragraphs"], s["tables"]], ensure_ascii=False))
        s["xml_hash"] = sha("".join(s.pop("_xml")))
    return {"type": "docx", "file": Path(path).name, "sections": sections}


def parse_xlsx(path):
    wb = load_workbook(str(path))  # формулы остаются формулами
    sections = []
    for ws in wb.worksheets:
        rows = [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
        while rows and not any(str(v).strip() for v in rows[-1]):
            rows.pop()
        h = sha(json.dumps(rows, ensure_ascii=False, default=str))
        sections.append({"title": ws.title, "level": 1, "path": ws.title, "dims": ws.dimensions,
                         "rows": rows, "text_hash": h, "xml_hash": h})
    return {"type": "xlsx", "file": Path(path).name, "sections": sections}


def parse(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return parse_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return parse_xlsx(path)
    raise SystemExit(f"Не поддерживается: {suffix} (только нативные .docx/.xlsx)")


def keyed(parsed):
    """path -> раздел; одинаковые пути получают суффикс #2, #3…"""
    out, seen = {}, {}
    for s in parsed["sections"]:
        n = seen[s["path"]] = seen.get(s["path"], 0) + 1
        out[s["path"] if n == 1 else f"{s['path']} #{n}"] = s
    return out


def compare(a, b, mode="xml_hash"):
    """Сравнение двух разобранных документов по разделам."""
    ka, kb = keyed(a), keyed(b)
    return {
        "added": [k for k in kb if k not in ka],
        "removed": [k for k in ka if k not in kb],
        "changed": [k for k in ka if k in kb and ka[k][mode] != kb[k][mode]],
        "same": [k for k in ka if k in kb and ka[k][mode] == kb[k][mode]],
    }


def section_lines(s):
    """Раздел как список строк — для человекочитаемого диффа и индексации."""
    lines = list(s.get("paragraphs", []))
    for t in s.get("tables", []):
        lines += [" | ".join(str(c) for c in row) for row in t]
    lines += [" | ".join(str(c) for c in row) for row in s.get("rows", [])]
    return lines
