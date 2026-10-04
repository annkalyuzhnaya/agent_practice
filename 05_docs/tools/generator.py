"""Генератор: сборка docx/xlsx по шаблону и точечное обновление больших документов.

  python tools/generator.py fill шаблон.docx данные.json --out отчёт.docx
  python tools/generator.py fill шаблон.xlsx данные.json --out сводка.xlsx
  python tools/generator.py update-section док.docx --section "Сроки" --text-file новый.md --out док_v2.docx
  python tools/generator.py set-cells книга.xlsx --sheet "Свод" --cells '{"B2": 15}' --out книга_v2.xlsx

Шаблон: {{ключ}} — подстановка значения; строка таблицы с {{rows:имя}} повторяется для каждого
элемента списка данные["имя"], внутри неё {{поле}} — поля элемента.
Ничего не выдумывает: незаполненные метки перечисляются, с --strict это ошибка (код 2).
update-section заменяет только абзацы указанного раздела; таблицы раздела и всё остальное не трогает.
"""
import argparse
import copy
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

import doclib

MARK = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
ROWS = re.compile(r"\{\{\s*rows:([^{}]+?)\s*\}\}")


def render(text, data, missing):
    def sub(m):
        key = m.group(1)
        if key in data and not isinstance(data[key], (list, dict)):
            return str(data[key])
        missing.add(key)
        return m.group(0)
    return MARK.sub(sub, text)


def set_text(p, text):
    """Меняет текст абзаца, сохраняя оформление первого фрагмента."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def fill_paragraphs(paragraphs, data, missing):
    for p in paragraphs:
        if "{{" in p.text:
            set_text(p, render(p.text, data, missing))


def fill_table(table, data, missing):
    for row in list(table.rows):
        m = ROWS.search("".join(c.text for c in row.cells))
        if not m:
            continue
        items = data.get(m.group(1))
        if not isinstance(items, list):
            missing.add(f"rows:{m.group(1)}")
            continue
        for item in items:
            tr = copy.deepcopy(row._tr)
            row._tr.addprevious(tr)
            for tc in tr.iterchildren():
                if tc.tag.endswith("}tc"):
                    for p_el in tc.iterchildren():
                        if p_el.tag.endswith("}p"):
                            p = Paragraph(p_el, table)
                            if "{{" in p.text:
                                set_text(p, render(ROWS.sub("", p.text), item, missing))
        row._tr.getparent().remove(row._tr)
    for row in table.rows:
        for cell in row.cells:
            fill_paragraphs(cell.paragraphs, data, missing)
            for nested in cell.tables:
                fill_table(nested, data, missing)


def fill_docx(template, data, out):
    doc = Document(str(template))
    missing = set()
    fill_paragraphs(doc.paragraphs, data, missing)
    for t in doc.tables:
        fill_table(t, data, missing)
    for s in doc.sections:
        for part in (s.header, s.footer):
            fill_paragraphs(part.paragraphs, data, missing)
            for t in part.tables:
                fill_table(t, data, missing)
    doc.save(str(out))
    return missing


def fill_xlsx(template, data, out):
    wb = load_workbook(str(template))
    missing = set()
    for ws in wb.worksheets:
        for row in list(ws.iter_rows()):
            hit = next((ROWS.search(c.value) for c in row if isinstance(c.value, str) and ROWS.search(c.value)), None)
            if not hit:
                continue
            items = data.get(hit.group(1))
            if not isinstance(items, list):
                missing.add(f"rows:{hit.group(1)}")
                continue
            r0 = row[0].row
            cells = [(c.column, ROWS.sub("", c.value) if isinstance(c.value, str) else c.value, copy.copy(c._style)) for c in row]
            if len(items) > 1:
                ws.insert_rows(r0 + 1, len(items) - 1)
            for i, item in enumerate(items or [{}]):
                for col, tpl, style in cells:
                    cell = ws.cell(row=r0 + i, column=col)
                    cell._style = copy.copy(style)
                    whole = MARK.fullmatch(tpl.strip()) if isinstance(tpl, str) else None
                    if whole and whole.group(1) in item:
                        cell.value = item[whole.group(1)]  # тип значения сохраняется (число остаётся числом)
                    else:
                        cell.value = render(tpl, item, missing) if isinstance(tpl, str) else tpl
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "{{" in c.value:
                    whole = MARK.fullmatch(c.value.strip())
                    if whole and whole.group(1) in data and not isinstance(data[whole.group(1)], (list, dict)):
                        c.value = data[whole.group(1)]
                    else:
                        c.value = render(c.value, data, missing)
    wb.save(str(out))
    return missing


def cmd_fill(args):
    data = json.loads(Path(args.data).read_text(encoding="utf-8"))
    fill = fill_docx if Path(args.template).suffix.lower() == ".docx" else fill_xlsx
    missing = fill(args.template, data, args.out)
    print(f"Собран документ: {args.out}")
    if missing:
        print("НЕ ЗАПОЛНЕНО (нет в данных): " + ", ".join(sorted(missing)))
        if args.strict:
            sys.exit(2)


def cmd_update_section(args):
    doc = Document(str(args.file))
    lines = [l.strip() for l in Path(args.text_file).read_text(encoding="utf-8").splitlines() if l.strip()]
    want = args.section.lower()
    blocks = list(doclib.iter_blocks(doc))
    heads = [i for i, b in enumerate(blocks) if isinstance(b, Paragraph) and doclib.heading_level(b)
             and want in b.text.lower()]
    if len(heads) != 1:
        found = [blocks[i].text for i in heads]
        sys.exit(f"Раздел «{args.section}»: найдено заголовков {len(heads)} {found} — уточните название")
    start = heads[0]
    old = []
    for b in blocks[start + 1:]:
        if isinstance(b, Paragraph):
            if doclib.heading_level(b):
                break
            if b._p.find(".//" + "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sectPr") is None:
                old.append(b)
    style = next((p.style for p in old if p.text.strip()), None)
    for p in old:
        p._p.getparent().remove(p._p)
    anchor = blocks[start]._p
    for line in lines:
        new_p = copy.deepcopy(anchor)
        for child in list(new_p):
            new_p.remove(child)
        anchor.addnext(new_p)
        para = Paragraph(new_p, blocks[start]._parent)
        para.add_run(line)
        para.style = style if style is not None else doc.styles["Normal"]
        anchor = new_p
    doc.save(str(args.out))
    print(f"Раздел «{blocks[start].text}» обновлён: было абзацев {len(old)}, стало {len(lines)}. Файл: {args.out}")
    print(f"Проверьте: python tools/integrity_diff.py \"{args.file}\" \"{args.out}\" --allow \"{blocks[start].text}\"")


def cmd_set_cells(args):
    wb = load_workbook(str(args.file))
    ws = wb[args.sheet]
    cells = json.loads(args.cells)
    for addr, value in cells.items():
        ws[addr] = value
    wb.save(str(args.out))
    print(f"Лист «{args.sheet}»: обновлено ячеек {len(cells)}. Файл: {args.out}")


def main():
    doclib.utf8()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("fill")
    p.add_argument("template")
    p.add_argument("data")
    p.add_argument("--out", required=True)
    p.add_argument("--strict", action="store_true")
    p.set_defaults(fn=cmd_fill)
    p = sub.add_parser("update-section")
    p.add_argument("file")
    p.add_argument("--section", required=True)
    p.add_argument("--text-file", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(fn=cmd_update_section)
    p = sub.add_parser("set-cells")
    p.add_argument("file")
    p.add_argument("--sheet", required=True)
    p.add_argument("--cells", required=True)
    p.add_argument("--out", required=True)
    p.set_defaults(fn=cmd_set_cells)
    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
